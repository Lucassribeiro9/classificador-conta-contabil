from pathlib import Path

import pytest
from openpyxl import Workbook

from core.plano_contas_parser import PlanoContasParseError, parse_plano_contas_xlsx


FIXTURES_DIR = Path(__file__).parent / "fixtures"


def _write_workbook(path, rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def test_parse_plano_contas_ignores_report_header_empty_rows_and_normalizes_accounts(
    tmp_path,
):
    xlsx_path = tmp_path / "plano-contas.xlsx"
    _write_workbook(
        xlsx_path,
        [
            ["Relatorio do Plano de Contas"],
            ["Empresa Modelo"],
            [],
            ["Codigo", "Tipo", "Classificacao", "Nome", "Grau"],
            [],
            [1000, "S", "1", "ATIVO", 1],
            [10046, "A", "1.1.01.01.02.10046", "BCO. SANTANDER", 6],
        ],
    )

    contas = parse_plano_contas_xlsx(xlsx_path)

    assert contas == [
        {
            "codigo": 1000,
            "tipo": "S",
            "classificacao": "1",
            "nome": "ATIVO",
            "grau": 1,
        },
        {
            "codigo": 10046,
            "tipo": "A",
            "classificacao": "1.1.01.01.02.10046",
            "nome": "BCO. SANTANDER",
            "grau": 6,
        },
    ]


def test_parse_plano_contas_reports_incomplete_rows_with_clear_message(tmp_path):
    xlsx_path = tmp_path / "plano-contas-incompleto.xlsx"
    _write_workbook(
        xlsx_path,
        [
            ["Codigo", "Tipo", "Classificacao", "Nome", "Grau"],
            [10046, "A", "1.1.01.01.02.10046", None, 6],
        ],
    )

    with pytest.raises(PlanoContasParseError, match="Linha 2.*nome"):
        parse_plano_contas_xlsx(xlsx_path)


def test_parse_plano_contas_ignores_system_footer_without_account_fields(tmp_path):
    xlsx_path = tmp_path / "plano-contas-com-rodape.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Codigo", "Tipo", "Classificacao", "Nome", "Grau"])
    sheet.append([10046, "A", "1.1.01.01.02.10046", "BANCO MODELO", 6])
    sheet.merge_cells("A3:E3")
    sheet["A3"] = "Relatorio emitido pelo sistema"
    workbook.save(xlsx_path)
    workbook.close()

    contas = parse_plano_contas_xlsx(xlsx_path)

    assert [conta["codigo"] for conta in contas] == [10046]


def test_parse_plano_contas_rejects_partial_account_in_merged_layout(tmp_path):
    xlsx_path = tmp_path / "plano-contas-parcial.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["B5"] = "Codigo"
    sheet["D5"] = "T"
    sheet["H5"] = "Classificacao"
    sheet["L5"] = "Nome"
    sheet["X5"] = "Grau"
    sheet.merge_cells("A6:B6")
    sheet["A6"] = 10046
    workbook.save(xlsx_path)
    workbook.close()

    with pytest.raises(PlanoContasParseError, match="Linha 6.*tipo"):
        parse_plano_contas_xlsx(xlsx_path)


def test_parse_plano_contas_accepts_dominio_layout_with_merged_cells():
    contas = parse_plano_contas_xlsx(
        FIXTURES_DIR / "plano_contas_dominio_sintetico.xlsx"
    )

    assert [conta["codigo"] for conta in contas] == [1000, 10046]
    assert contas[1] == {
        "codigo": 10046,
        "tipo": "A",
        "classificacao": "1.1.01.01.02.10046",
        "nome": "BANCO MODELO",
        "grau": 6,
    }
